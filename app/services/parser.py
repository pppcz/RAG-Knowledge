"""文档解析：PDF / Word(.docx) / Markdown / TXT → 纯文本。"""

import io
import logging
import re

from app.core.errors import AppError

logger = logging.getLogger(__name__)

# 常见中文编码依次尝试
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1")

HTML_TAG = re.compile(r"<[^>]+>")
MULTI_BLANK = re.compile(r"\n{3,}")
TRAILING_SPACE = re.compile(r"[ \t]+(?=\n)")


def normalize_text(text: str) -> str:
    """统一换行、压缩空行与行尾空白。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = TRAILING_SPACE.sub("", text)
    text = MULTI_BLANK.sub("\n\n", text)
    return text.strip()


def decode_text(data: bytes) -> str:
    """按常见编码解码文本文件。"""
    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    """提取 PDF 文本；扫描件（无文本层）视为不可解析。"""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        logger.warning("PDF 解析失败: %s", exc)
        raise AppError(42201, "PDF 文件解析失败，文件可能已损坏", 422) from exc

    text = "\n\n".join(part for part in pages if part.strip())
    if not text.strip():
        raise AppError(42201, "PDF 中没有可提取的文本（可能是扫描件）", 422)
    return text


def _extract_docx(data: bytes) -> str:
    """提取 Word 段落与表格文本。"""
    from docx import Document as DocxDocument

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        logger.warning("DOCX 解析失败: %s", exc)
        raise AppError(42201, "Word 文件解析失败，文件可能已损坏", 422) from exc

    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(part for part in parts if part.strip())
    if not text.strip():
        raise AppError(42201, "Word 文件中没有可提取的文本", 422)
    return text


def _extract_markdown(data: bytes) -> str:
    """Markdown：渲染为 HTML 后去标签，保留正文结构。"""
    import markdown as md

    raw = decode_text(data)
    html = md.markdown(raw, extensions=["tables", "fenced_code"])
    text = HTML_TAG.sub("", html)
    return text


def _extract_txt(data: bytes) -> str:
    """纯文本：按常见编码解码。"""
    return decode_text(data)


_EXTRACTORS = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "md": _extract_markdown,
    "txt": _extract_txt,
}


def extract_text(file_ext: str, data: bytes) -> str:
    """按扩展名提取纯文本，失败抛出 42201。"""
    extractor = _EXTRACTORS.get(file_ext.lower())
    if extractor is None:
        raise AppError(41501, f"不支持的文件类型：{file_ext}", 415)

    text = normalize_text(extractor(data))
    if not text:
        raise AppError(42201, "文档内容为空，无法索引", 422)
    return text
