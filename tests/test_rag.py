"""阶段 3：解析、切片、Embedding 与向量检索测试。"""

import io
from pathlib import Path

import pytest
from app.core.errors import AppError
from app.models.document import Chunk
from app.services import chunking, embeddings, parser
from app.services.embeddings import HashEmbeddingProvider
from httpx import AsyncClient
from sqlalchemy import func, select

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "docs"


def _cosine(left: list[float], right: list[float]) -> float:
    """两个已归一化向量的余弦相似度。"""
    return sum(a * b for a, b in zip(left, right, strict=True))


# ---------------------------------------------------------------- 解析器


def test_parser_txt_utf8() -> None:
    """TXT：UTF-8 解码并保留正文。"""
    text = parser.extract_text("txt", "报销流程：提交、审批、付款。".encode())

    assert "报销流程" in text


def test_parser_txt_gbk() -> None:
    """TXT：GBK 编码自动识别。"""
    text = parser.extract_text("txt", "差旅住宿标准".encode("gbk"))

    assert "差旅住宿标准" in text


def test_parser_markdown_strips_syntax() -> None:
    """Markdown：去掉语法标记，保留可见文本。"""
    raw = "# 标题\n\n- 第一项\n- 第二项\n".encode()

    text = parser.extract_text("md", raw)

    assert "标题" in text
    assert "第一项" in text
    assert "#" not in text


def test_parser_pdf_sample() -> None:
    """PDF：从样例 PDF 提取文本。"""
    data = (SAMPLE_DIR / "sample.pdf").read_bytes()

    text = parser.extract_text("pdf", data)

    assert "Reimbursement Process" in text


def test_parser_docx() -> None:
    """DOCX：段落与表格文本都能提取。"""
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("员工手册第一章：考勤制度")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "年假天数"
    table.rows[0].cells[1].text = "5 天"
    buffer = io.BytesIO()
    document.save(buffer)

    text = parser.extract_text("docx", buffer.getvalue())

    assert "考勤制度" in text
    assert "年假天数" in text
    assert "5 天" in text


def test_parser_invalid_pdf_raises() -> None:
    """异常路径：损坏的 PDF 抛出 42201。"""
    with pytest.raises(AppError) as excinfo:
        parser.extract_text("pdf", b"not a real pdf")

    assert excinfo.value.code == 42201
    assert excinfo.value.status_code == 422


def test_parser_unsupported_extension() -> None:
    """异常路径：不支持的扩展名抛出 41501。"""
    with pytest.raises(AppError) as excinfo:
        parser.extract_text("exe", b"MZ")

    assert excinfo.value.code == 41501


# ---------------------------------------------------------------- 切片


def test_split_text_respects_size_and_overlap() -> None:
    """切片长度不超过 CHUNK_SIZE，且相邻片段有重叠。"""
    text = "报销流程说明。" * 200

    chunks = chunking.split_text(text, chunk_size=100, chunk_overlap=20)

    assert len(chunks) > 1
    assert all(len(chunk) <= 100 for chunk in chunks)
    assert all(chunk.strip() for chunk in chunks)

    tail = chunks[0][-20:]
    assert any(tail[-10:] in chunk for chunk in chunks[1:])


def test_split_text_prefers_sentence_boundary() -> None:
    """切片优先在句末断开。"""
    text = "第一句话内容比较长。" * 30

    chunks = chunking.split_text(text, chunk_size=60, chunk_overlap=10)

    assert chunks[0].endswith("。")


def test_split_text_empty() -> None:
    """空文本返回空列表。"""
    assert chunking.split_text("   \n\n ", 100, 10) == []


def test_split_text_invalid_size() -> None:
    """异常路径：CHUNK_SIZE <= 0 抛错。"""
    with pytest.raises(AppError):
        chunking.split_text("abc", 0, 0)


# ---------------------------------------------------------------- Embedding


async def test_hash_embedding_shape_and_norm() -> None:
    """hash 向量：维度正确、L2 归一化、取值确定。"""
    embedder = HashEmbeddingProvider(dim=1536)

    [vector] = await embedder.embed_texts(["报销流程与审批规则"])

    assert len(vector) == 1536
    assert pytest.approx(sum(value * value for value in vector), rel=1e-6) == 1.0


async def test_hash_embedding_is_deterministic() -> None:
    """同一文本（含跨实例）向量完全一致。"""
    first = await HashEmbeddingProvider(dim=256).embed_texts(["差旅住宿标准"])
    second = await HashEmbeddingProvider(dim=256).embed_texts(["差旅住宿标准"])

    assert first == second


async def test_hash_embedding_ranks_similar_text_higher() -> None:
    """hash 向量只捕捉词汇重叠：相近文本相似度更高。"""
    embedder = HashEmbeddingProvider(dim=1536)
    query = "报销流程 审批 发票"
    [query_vector, similar, different] = await embedder.embed_texts(
        [query, "报销流程需要提交发票，主管审批后付款", "服务器机房温度与网络带宽巡检记录"]
    )

    assert _cosine(query_vector, similar) > _cosine(query_vector, different)


async def test_get_embedder_defaults_to_hash(monkeypatch: pytest.MonkeyPatch) -> None:
    """默认配置返回 hash Provider。"""
    embedder = embeddings.get_embedder()

    assert isinstance(embedder, HashEmbeddingProvider)


# ---------------------------------------------------------------- 接口


async def _create_kb(client: AsyncClient, headers: dict[str, str], name: str) -> dict:
    """创建知识库。"""
    resp = await client.post("/api/kb", json={"name": name}, headers=headers)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def _upload(
    client: AsyncClient,
    headers: dict[str, str],
    filename: str,
    content: bytes,
    kb_id: int | None = None,
):
    """上传文档。"""
    data = {"kb_id": str(kb_id)} if kb_id is not None else {}
    return await client.post(
        "/api/docs/upload",
        files={"file": (filename, content, "text/plain")},
        data=data,
        headers=headers,
    )


async def test_upload_indexes_document(client: AsyncClient, tenant_admin: dict, db_session) -> None:
    """上传后立即解析入库：status=ready，切片写入 chunks 且带向量。"""
    kb = await _create_kb(client, tenant_admin["headers"], "索引库")

    resp = await _upload(
        client,
        tenant_admin["headers"],
        "sample.md",
        (SAMPLE_DIR / "sample.md").read_bytes(),
        kb_id=kb["id"],
    )

    assert resp.status_code == 201, resp.text
    document = resp.json()
    assert document["status"] == "ready"
    assert document["chunk_count"] >= 1
    assert document["error_message"] is None

    total = await db_session.scalar(
        select(func.count()).select_from(Chunk).where(Chunk.document_id == document["id"])
    )
    with_embedding = await db_session.scalar(
        select(func.count())
        .select_from(Chunk)
        .where(Chunk.document_id == document["id"], Chunk.embedding.is_not(None))
    )
    assert total == document["chunk_count"]
    assert with_embedding == total


async def test_search_returns_relevant_chunk(client: AsyncClient, tenant_admin: dict) -> None:
    """检索：查询「报销流程」命中报销文档。"""
    kb = await _create_kb(client, tenant_admin["headers"], "检索库")
    await _upload(
        client,
        tenant_admin["headers"],
        "sample.md",
        (SAMPLE_DIR / "sample.md").read_bytes(),
        kb_id=kb["id"],
    )

    resp = await client.get(
        "/api/kb/search",
        params={"kb_id": kb["id"], "q": "报销流程", "top_k": 3},
        headers=tenant_admin["headers"],
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["total"] >= 1
    assert body["kb_id"] == kb["id"]
    top = body["items"][0]
    assert "报销" in top["content"]
    assert top["score"] > 0
    assert top["filename"] == "sample.md"


async def test_search_empty_kb_returns_no_hits(client: AsyncClient, tenant_admin: dict) -> None:
    """空知识库返回空结果而非报错。"""
    kb = await _create_kb(client, tenant_admin["headers"], "空库")

    resp = await client.get(
        "/api/kb/search",
        params={"kb_id": kb["id"], "q": "报销流程"},
        headers=tenant_admin["headers"],
    )

    assert resp.status_code == 200
    assert resp.json()["total"] == 0
    assert resp.json()["items"] == []


async def test_search_isolated_between_kbs(client: AsyncClient, tenant_admin: dict) -> None:
    """检索只在指定知识库内进行。"""
    kb_a = await _create_kb(client, tenant_admin["headers"], "库A")
    kb_b = await _create_kb(client, tenant_admin["headers"], "库B")
    await _upload(
        client,
        tenant_admin["headers"],
        "sample.md",
        (SAMPLE_DIR / "sample.md").read_bytes(),
        kb_id=kb_a["id"],
    )

    resp = await client.get(
        "/api/kb/search",
        params={"kb_id": kb_b["id"], "q": "报销流程"},
        headers=tenant_admin["headers"],
    )

    assert resp.status_code == 200
    assert resp.json()["total"] == 0


async def test_search_foreign_tenant_kb_404(client: AsyncClient, tenant_admin: dict) -> None:
    """异常路径：检索其它租户的知识库返回 404。"""
    other = await client.post(
        "/api/auth/register",
        json={
            "tenant_name": "Other",
            "tenant_slug": "other",
            "username": "admin",
            "email": "other@example.com",
            "password": "secret123",
        },
    )
    other_headers = {"Authorization": f"Bearer {other.json()['access_token']}"}
    other_kb = await _create_kb(client, other_headers, "别的库")

    resp = await client.get(
        "/api/kb/search",
        params={"kb_id": other_kb["id"], "q": "报销流程"},
        headers=tenant_admin["headers"],
    )

    assert resp.status_code == 404
    assert resp.json()["code"] == 40401


async def test_search_validation(client: AsyncClient, tenant_admin: dict) -> None:
    """异常路径：缺少 q、top_k 越界都返回 40001。"""
    kb = await _create_kb(client, tenant_admin["headers"], "校验库")

    missing_q = await client.get(
        "/api/kb/search",
        params={"kb_id": kb["id"]},
        headers=tenant_admin["headers"],
    )
    assert missing_q.status_code == 422
    assert missing_q.json()["code"] == 40001

    bad_top_k = await client.get(
        "/api/kb/search",
        params={"kb_id": kb["id"], "q": "x", "top_k": 0},
        headers=tenant_admin["headers"],
    )
    assert bad_top_k.status_code == 422


async def test_reindex_rebuilds_chunks(client: AsyncClient, tenant_admin: dict) -> None:
    """重建索引：切片数量保持一致，状态 ready。"""
    kb = await _create_kb(client, tenant_admin["headers"], "重建库")
    uploaded = await _upload(
        client,
        tenant_admin["headers"],
        "sample.txt",
        (SAMPLE_DIR / "sample.txt").read_bytes(),
        kb_id=kb["id"],
    )
    document_id = uploaded.json()["id"]
    first_count = uploaded.json()["chunk_count"]

    resp = await client.post(f"/api/docs/{document_id}/reindex", headers=tenant_admin["headers"])

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["status"] == "ready"
    assert body["chunk_count"] == first_count


async def test_reindex_marks_failed_for_broken_file(
    client: AsyncClient, tenant_admin: dict
) -> None:
    """异常路径：无法解析的 PDF 标记 failed 并给出原因（接口仍返回 200）。"""
    kb = await _create_kb(client, tenant_admin["headers"], "坏文件库")
    uploaded = await _upload(
        client,
        tenant_admin["headers"],
        "broken.pdf",
        b"%PDF-1.4 this is not a valid pdf",
        kb_id=kb["id"],
    )

    assert uploaded.status_code == 201
    document = uploaded.json()
    assert document["status"] == "failed"
    assert document["error_message"]
    assert document["chunk_count"] == 0

    reindexed = await client.post(
        f"/api/docs/{document['id']}/reindex", headers=tenant_admin["headers"]
    )
    assert reindexed.status_code == 200
    assert reindexed.json()["status"] == "failed"


async def test_member_can_search_but_not_reindex(
    client: AsyncClient, tenant_admin: dict, member_headers: dict
) -> None:
    """RBAC：member 有 kb:read 可检索，但没有 docs:write 不能重建索引。"""
    kb = await _create_kb(client, tenant_admin["headers"], "权限库")
    uploaded = await _upload(
        client,
        tenant_admin["headers"],
        "sample.md",
        (SAMPLE_DIR / "sample.md").read_bytes(),
        kb_id=kb["id"],
    )

    search = await client.get(
        "/api/kb/search",
        params={"kb_id": kb["id"], "q": "报销流程"},
        headers=member_headers,
    )
    assert search.status_code == 200
    assert search.json()["total"] >= 1

    reindex = await client.post(
        f"/api/docs/{uploaded.json()['id']}/reindex", headers=member_headers
    )
    assert reindex.status_code == 403
    assert reindex.json()["detail"]["required_permission"] == "docs:write"


async def test_deleted_document_excluded_from_search(
    client: AsyncClient, tenant_admin: dict
) -> None:
    """删除文档后不再出现在检索结果中。"""
    kb = await _create_kb(client, tenant_admin["headers"], "删除库")
    uploaded = await _upload(
        client,
        tenant_admin["headers"],
        "sample.md",
        (SAMPLE_DIR / "sample.md").read_bytes(),
        kb_id=kb["id"],
    )
    document_id = uploaded.json()["id"]

    await client.delete(f"/api/docs/{document_id}", headers=tenant_admin["headers"])

    resp = await client.get(
        "/api/kb/search",
        params={"kb_id": kb["id"], "q": "报销流程"},
        headers=tenant_admin["headers"],
    )

    assert resp.status_code == 200
    assert resp.json()["total"] == 0
